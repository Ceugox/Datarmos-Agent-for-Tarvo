from pathlib import Path
from collections import Counter, defaultdict
from decimal import Decimal, ROUND_HALF_UP, ROUND_DOWN
from datetime import datetime
import csv
import json
import re
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCES = ROOT / 'sources'
OUT = ROOT / 'analysis'

def money(value):
    return Decimal(str(value or 0).replace('.', '').replace(',', '.'))

def dec(value):
    return Decimal(str(value or 0))

def rounded(value):
    return value.quantize(Decimal('.01'), rounding=ROUND_HALF_UP)

def total(rows, field):
    return sum((r[field] for r in rows), Decimal(0))

def save(name, rows):
    (OUT / name).write_text(json.dumps(rows, ensure_ascii=False, indent=2, default=str), encoding='utf-8')

with (SOURCES / '01_orbe_relatorio_comissoes_set2026.csv').open(encoding='utf-8-sig') as stream:
    checkout = list(csv.DictReader(stream, delimiter=';'))
for number, row in enumerate(checkout, 2):
    row['source_line'] = number
    for key in ['valor_bruto', 'taxa_aplicada_pct', 'valor_taxa', 'valor_liquido']:
        row[key] = money(row[key])
    row['qtd_parcelas'] = int(row['qtd_parcelas']) if row['qtd_parcelas'] else None

wb = load_workbook(SOURCES / '02_tabela_taxas_contratadas.xlsx', data_only=True)
contracts = []
for number, cells in enumerate(wb['Taxas vigentes'].values, 1):
    if number < 6 or not cells[1]:
        continue
    row = dict(zip(['name', 'cnpj', 'ec', 'debit', 'credit', 'installment_short', 'installment_long', 'pix', 'advance', 'spread', 'validity', 'notes'], cells))
    row['source_row'] = number
    row['cnpj'] = re.sub(r'\D', '', str(row['cnpj'])).zfill(14)
    for key in ['debit', 'credit', 'installment_short', 'installment_long', 'pix', 'advance', 'spread']:
        value = row[key]
        if value is not None and value != 'padrão':
            row[key] = dec(value) if not isinstance(value, str) else money(value.replace('%', '')) / 100
    contracts.append(row)

by_ec = defaultdict(list)
for row in contracts:
    if row['ec']:
        by_ec[str(row['ec'])].append(row)

comparisons = []
unmapped = []
for row in checkout:
    choices = by_ec[row['cod_estabelecimento']]
    if not choices:
        unmapped.append(row)
        continue
    if row['tipo_lancamento'] != 'VENDA':
        continue
    contract = choices[0]
    if len(choices) > 1:
        date = datetime.strptime(row['data_movimento'], '%d/%m/%Y')
        contract = choices[0 if date.day < 15 else 1]
    modality = row['modalidade']
    key = {'DEBITO': 'debit', 'CREDITO A VISTA': 'credit', 'PIX': 'pix'}.get(modality)
    if modality == 'CREDITO PARCELADO':
        key = 'installment_short' if row['qtd_parcelas'] <= 6 else 'installment_long'
    if not key or contract[key] is None:
        continue
    expected = (row['valor_bruto'] * contract[key]).quantize(Decimal('.01'), rounding=ROUND_DOWN)
    comparisons.append({**row, 'contract_name': contract['name'], 'cnpj': contract['cnpj'], 'contract_row': contract['source_row'], 'contract_pct': contract[key] * 100, 'expected_fee': expected, 'difference': row['valor_taxa'] - expected})

fx_wb = load_workbook(SOURCES / 'PonteCambio_Relatorio_TARVO_092026.xlsx', data_only=True)
fx = []
keys = ['operation', 'contracted', 'settled', 'cnpj', 'name', 'currency', 'foreign_amount', 'client_rate', 'reference_rate', 'brl', 'spread', 'spread_pct', 'estimated_spread', 'iof', 'status', 'notes']
for number, cells in enumerate(fx_wb['Operações'].values, 1):
    if number < 6 or not str(cells[0]).startswith('PC-'):
        continue
    row = dict(zip(keys, cells))
    row['source_row'] = number
    for key in ['foreign_amount', 'client_rate', 'reference_rate', 'brl', 'spread', 'spread_pct', 'estimated_spread', 'iof']:
        row[key] = dec(row[key])
    row['cnpj'] = re.sub(r'\D', '', str(row['cnpj'])).zfill(14)
    fx.append(row)
groups = defaultdict(list)
for row in fx:
    groups[row['operation']].append(row)
duplicates = {key: rows for key, rows in groups.items() if len(rows) > 1}
unique_fx = [rows[0] for rows in groups.values()]
eligible_fx = [r for r in unique_fx if r['settled'] and r['settled'].month == 9 and r['status'] in ['APROVADO', 'APROVADO C/ RESSALVA']]

aurora = []
lines = (SOURCES / '05a_extrato_banco_aurora_cc4471_set2026.csv').read_text(encoding='utf-8-sig').splitlines()
for number, line in enumerate(lines, 1):
    cells = next(csv.reader([line], delimiter=';'))
    if len(cells) == 5 and re.match(r'\d{2}/\d{2}/2026', cells[0]) and cells[3]:
        aurora.append({'ref': f'L{number}', 'date': cells[0], 'description': cells[1], 'amount': money(cells[3]), 'balance': money(cells[4])})
ofx = (SOURCES / '05b_extrato_banco_meridiano_cc90233_set2026.ofx').read_text(encoding='cp1252')
meridiano = []
for block in re.findall(r'<STMTTRN>(.*?)</STMTTRN>', ofx, re.S):
    tags = dict(re.findall(r'<(\w+)>([^<\r\n]+)', block))
    meridiano.append({'ref': tags['FITID'], 'date': tags['DTPOSTED'][:8], 'description': tags.get('MEMO', ''), 'amount': dec(tags['TRNAMT'])})

summary = {
    'checkout': {'count': len(checkout), 'unique_ids': len(set(r['id_transacao'] for r in checkout)), 'ecs': len(set(r['cod_estabelecimento'] for r in checkout)),
        'types': {t: {'count': len(rows := [r for r in checkout if r['tipo_lancamento'] == t]), 'gross': total(rows, 'valor_bruto'), 'fee': total(rows, 'valor_taxa'), 'net': total(rows, 'valor_liquido')} for t in sorted(set(r['tipo_lancamento'] for r in checkout))},
        'modalities': dict(Counter(r['modalidade'] for r in checkout)),
        'gross': total(checkout, 'valor_bruto'), 'fee': total(checkout, 'valor_taxa'), 'net': total(checkout, 'valor_liquido'),
        'arithmetic_failures': [r for r in checkout if r['valor_bruto'] - r['valor_taxa'] != r['valor_liquido']],
        'unmapped': {ec: {'count': len(rows := [r for r in unmapped if r['cod_estabelecimento'] == ec]), 'name': rows[0]['nome_estabelecimento'], 'gross': total(rows, 'valor_bruto'), 'fee': total(rows, 'valor_taxa')} for ec in sorted(set(r['cod_estabelecimento'] for r in unmapped))},
        'compared_count': len(comparisons),
        'rate_mismatch_count': sum(r['contract_pct'] != r['taxa_aplicada_pct'] for r in comparisons),
        'positive_difference': sum((r['difference'] for r in comparisons if r['difference'] > 0), Decimal(0)),
        'negative_difference': sum((r['difference'] for r in comparisons if r['difference'] < 0), Decimal(0))},
    'contracts': {'rows': len(contracts), 'cnpjs': len(set(r['cnpj'] for r in contracts)), 'missing_ec': [r for r in contracts if not r['ec']]},
    'fx': {'rows': len(fx), 'unique': len(unique_fx), 'cnpjs': len(set(r['cnpj'] for r in fx)), 'duplicates': duplicates,
        'status': {status: {'count': len(rows := [r for r in unique_fx if r['status'] == status]), 'brl': total(rows,'brl'), 'spread': total(rows,'spread')} for status in sorted(set(r['status'] for r in fx))},
        'raw_spread': total(fx,'spread'), 'unique_spread': total(unique_fx,'spread'), 'raw_brl': total(fx,'brl'), 'unique_brl': total(unique_fx,'brl'),
        'eligible_count': len(eligible_fx), 'eligible_brl': total(eligible_fx,'brl'), 'eligible_spread': total(eligible_fx,'spread'),
        'eligible_half_rounded_each': sum((rounded(r['spread']/2) for r in eligible_fx), Decimal(0)),
        'negative': [r for r in unique_fx if r['spread'] < 0], 'not_settled_september': [r for r in unique_fx if not r['settled'] or r['settled'].month != 9],
        'summary_rows': list(fx_wb['Resumo'].values)},
    'banks': {'aurora_count': len(aurora), 'aurora_opening': money('612.334,81'), 'aurora_movement': total(aurora, 'amount'), 'aurora_closing': aurora[-1]['balance'],
        'meridiano_count': len(meridiano), 'meridiano_movement': total(meridiano,'amount'), 'meridiano_closing': dec(re.search(r'<BALAMT>([^<\r\n]+)',ofx)[1]),
        'meridiano_by_description': {d: {'count': len(rows := [r for r in meridiano if r['description'] == d]), 'amount': total(rows,'amount')} for d in sorted(set(r['description'] for r in meridiano))}}
}
save('findings.json', summary)
save('checkout_comparisons.json', comparisons)
save('fx_operations.json', fx)
save('bank_entries.json', {'aurora': aurora, 'meridiano': meridiano})
save('contracts.json', contracts)
if __name__ == '__main__':
    print(json.dumps(summary, ensure_ascii=True, indent=2, default=str))
