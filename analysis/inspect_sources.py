from pathlib import Path
from email import policy
from email.parser import BytesParser
from openpyxl import load_workbook
from pypdf import PdfReader
import json
root=Path('.')
(root/'analysis'/'instructions.txt').write_text('\n\n'.join(f'Página {i+1}\n'+p.extract_text() for i,p in enumerate(PdfReader('01_instrucoes_participantes.docx.pdf').pages)),encoding='utf-8')
message=BytesParser(policy=policy.default).parsebytes((root/'sources'/'03_email_ponte_cambio_relatorio_set2026.eml').read_bytes())
for part in message.walk():
 if part.get_filename():
  (root/'sources'/part.get_filename()).write_bytes(part.get_payload(decode=True))
 elif part.get_content_type()=='text/plain':
  print(part.get_content())
for path in (root/'sources').glob('*.xlsx'):
 wb=load_workbook(path,data_only=False)
 print('\nWORKBOOK',path.name)
 for ws in wb:
  print('SHEET',ws.title,ws.max_row,ws.max_column)
  for row in list(ws.values)[:12]:print(row)
  (root/'analysis'/(path.stem+'_'+ws.title.replace('/','-')+'.json')).write_text(json.dumps(list(ws.values),ensure_ascii=False,default=str,indent=2),encoding='utf-8')
