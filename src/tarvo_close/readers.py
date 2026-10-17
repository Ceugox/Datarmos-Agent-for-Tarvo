from datetime import datetime
from email import policy
from email.parser import BytesParser
from io import BytesIO, StringIO
from pathlib import Path
from tempfile import TemporaryDirectory
import csv
import hashlib
import re
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook

from .domain import ClosingInputs, brazilian_decimal, cnpj, decimal, normalize, parse_rate, parse_validity


def decode(content):
    try:
        return content.decode('utf-8-sig')
    except UnicodeDecodeError:
        return content.decode('cp1252')


def reference(filename, key, row, sheet=''):
    return {'arquivo_origem': filename, 'ref_origem': key, 'linha_origem': row, 'aba_origem': sheet}


def read_checkout(path):
    reader = csv.DictReader(StringIO(decode(path.read_bytes())), delimiter=';')
    expected = {'id_transacao', 'data_movimento', 'cod_estabelecimento', 'nome_estabelecimento', 'tipo_lancamento', 'modalidade', 'qtd_parcelas', 'bandeira', 'valor_bruto', 'taxa_aplicada_pct', 'valor_taxa', 'valor_liquido', 'id_referencia', 'data_venda_original'}
    if set(reader.fieldnames or []) != expected:
        raise ValueError('Unexpected checkout columns')
    rows = []
    for row in reader:
        if None in row or any(value is None for value in row.values()):
            raise ValueError(f'Invalid checkout row at line {reader.line_num}')
        row['_source'] = reference(path.name, row['id_transacao'], reader.line_num)
        row['date'] = datetime.strptime(row['data_movimento'], '%d/%m/%Y').date()
        for key in ['valor_bruto', 'taxa_aplicada_pct', 'valor_taxa', 'valor_liquido']:
            row[key] = brazilian_decimal(row[key])
        if row['valor_bruto'] - row['valor_taxa'] != row['valor_liquido']:
            raise ValueError(f'Checkout arithmetic mismatch: {row["id_transacao"]}')
        row['qtd_parcelas'] = int(row['qtd_parcelas']) if row['qtd_parcelas'] else None
        if row['tipo_lancamento'] not in {'VENDA', 'ESTORNO', 'ANTECIPACAO', 'REVERSAO ANTECIPACAO'}:
            raise ValueError(f'Unknown checkout event: {row["tipo_lancamento"]}')
        rows.append(row)
    return rows


def read_contracts(path):
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if 'Taxas vigentes' not in workbook.sheetnames:
            raise ValueError('Missing current contracts sheet')
        rows = []
        found_header = False
        for number, cells in enumerate(workbook['Taxas vigentes'].values, 1):
            if normalize(cells[0]) == 'LOJISTA' and normalize(cells[1]) == 'CNPJ':
                found_header = True
                continue
            if not found_header or not any(v is not None for v in cells):
                continue
            if cells[0] and str(cells[0]).startswith('*'):
                continue
            if not cells[1]:
                raise ValueError(f'Missing contract identity at row {number}')
            if 'INATIVO' in normalize(cells[11]):
                continue
            start, end = parse_validity(cells[10])
            row = {'name': cells[0], 'cnpj': cnpj(cells[1]), 'ec': str(cells[2]).strip() if cells[2] is not None else None, 'start': start, 'end': end, 'notes': cells[11] or ''}
            for index, key in enumerate(['debit', 'credit', 'installment_short', 'installment_long', 'pix', 'advance', 'spread'], 3):
                row[key] = parse_rate(cells[index])
            row['_source'] = reference(path.name, f'Taxas vigentes!L{number}', number, 'Taxas vigentes')
            rows.append(row)
        if not found_header or not rows:
            raise ValueError('No current contracts found')
        return rows
    finally:
        workbook.close()


def read_remittance_email(path):
    message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
    body = '\n'.join(part.get_content() for part in message.walk() if part.get_content_type() == 'text/plain' and not part.get_filename())
    attachments = [(part.get_filename(), part.get_payload(decode=True)) for part in message.walk() if part.get_filename() and part.get_filename().lower().endswith('.xlsx')]
    if len(attachments) != 1:
        raise ValueError('Expected exactly one remittance workbook attachment')
    filename, content = attachments[0]
    if Path(filename).name != filename or '/' in filename or '\\' in filename:
        raise ValueError('Attachment filename contains a path')
    workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    rows = []
    try:
        if 'Operações' not in workbook.sheetnames:
            raise ValueError('Missing remittance operations sheet')
        found_header = False
        for number, cells in enumerate(workbook['Operações'].values, 1):
            if normalize(cells[0]) in {'NO OPERACAO', 'N OPERACAO'}:
                found_header = True
                continue
            if not found_header or not any(v is not None for v in cells):
                continue
            if cells[0] in {'Total de operações:', 'Spread total R$:'}:
                continue
            if not cells[0] or not isinstance(cells[1], datetime):
                raise ValueError(f'Unrecognized remittance row {number}')
            row = {'operation': str(cells[0]), 'contracted': cells[1].date(), 'settled': cells[2].date() if isinstance(cells[2], datetime) else None, 'cnpj': cnpj(cells[3]), 'name': str(cells[4]), 'currency': str(cells[5]), 'status': str(cells[14]), 'notes': cells[15] or ''}
            for index, key in [(6,'foreign_amount'),(7,'client_rate'),(8,'reference_rate'),(9,'brl'),(10,'spread'),(11,'spread_pct'),(12,'estimated_spread'),(13,'iof')]:
                row[key] = decimal(cells[index])
            row['_source'] = reference(filename, row['operation'], number, 'Operações')
            rows.append(row)
        if not found_header or not rows:
            raise ValueError('No remittance operations found')
    finally:
        workbook.close()
    return rows, body, {filename: content}


def read_aurora(path):
    reader = csv.reader(StringIO(decode(path.read_bytes())), delimiter=';')
    rows = []
    opening = None
    balance = None
    closing = None
    for cells in reader:
        if len(cells) != 5:
            raise ValueError(f'Unexpected bank row shape: L{reader.line_num}')
        if 'SALDO ANTERIOR' in normalize(cells[1]):
            opening = balance = brazilian_decimal(cells[4])
            continue
        if normalize(cells[1]) == 'SALDO FINAL':
            closing = brazilian_decimal(cells[4])
            continue
        if not re.fullmatch(r'\d{2}/\d{2}/\d{4}', cells[0]):
            continue
        amount = brazilian_decimal(cells[3])
        reported_balance = brazilian_decimal(cells[4])
        if balance is None or balance + amount != reported_balance:
            raise ValueError(f'Bank running balance mismatch: L{reader.line_num}')
        balance = reported_balance
        rows.append({'bank':'AURORA', 'date':datetime.strptime(cells[0],'%d/%m/%Y').date(), 'description':cells[1], 'amount':amount, '_source':reference(path.name,f'L{reader.line_num}',reader.line_num)})
    if not rows or closing is None or balance != closing:
        raise ValueError('Missing bank rows or closing balance mismatch')
    return rows, {'opening': opening, 'closing': closing, 'opening_status': 'EXPLICITO'}


def read_meridiano(path):
    content = decode(path.read_bytes())
    rows = []
    for match in re.finditer(r'<STMTTRN>(.*?)</STMTTRN>', content, re.S):
        tags = dict(re.findall(r'<(\w+)>([^<\r\n]+)', match.group(1)))
        for key in ['FITID','DTPOSTED','TRNAMT','MEMO']:
            if not tags.get(key):
                raise ValueError(f'Missing OFX field: {key}')
        rows.append({'bank':'MERIDIANO', 'date':datetime.strptime(tags['DTPOSTED'][:8],'%Y%m%d').date(), 'description':tags['MEMO'], 'amount':decimal(tags['TRNAMT']), '_source':reference(path.name,tags['FITID'],content[:match.start()].count('\n')+1)})
    closing_match = re.search(r'<BALAMT>([^<\r\n]+)', content)
    if not rows or not closing_match:
        raise ValueError('Missing OFX transactions or ledger balance')
    closing = decimal(closing_match[1])
    opening = closing - sum((r['amount'] for r in rows), decimal('0'))
    return rows, {'opening':opening, 'closing':closing, 'opening_status':'INFERIDO'}


SOURCE_PATTERNS = [('checkout','01_*.csv'),('contracts','02_*.xlsx'),('email','03_*.eml'),('accounts','04_*.csv'),('aurora','05a_*.csv'),('meridiano','05b_*.ofx')]


def read_accounts(path):
    accounts = csv.DictReader(StringIO(decode(path.read_bytes())), delimiter=';')
    expected = {'COD_CONTA','DESCRICAO','TIPO','NATUREZA','ACEITA_LANCTO','COD_SUPERIOR','DT_CRIACAO'}
    if set(accounts.fieldnames or []) != expected:
        raise ValueError('Unexpected chart of accounts columns')
    result = {}
    for row in accounts:
        if row['COD_CONTA'] in result:
            raise ValueError('Duplicate chart account')
        result[row['COD_CONTA']] = row
    return result


def load_source(bundle, key, path):
    if key == 'checkout':
        bundle.checkout = read_checkout(path)
    elif key == 'contracts':
        bundle.contracts = read_contracts(path)
    elif key == 'email':
        remittances, body, attachments = read_remittance_email(path)
        bundle.remittances, bundle.email_body, bundle.attachments = remittances, body, attachments
        for name, content in attachments.items():
            bundle.sources[name] = {'sha256':hashlib.sha256(content).hexdigest(), 'parent':path.name, 'kind':'attachment'}
    elif key == 'accounts':
        bundle.accounts = read_accounts(path)
    else:
        bank = 'AURORA' if key == 'aurora' else 'MERIDIANO'
        rows, balances = (read_aurora if key == 'aurora' else read_meridiano)(path)
        bundle.bank_entries.extend(rows)
        bundle.bank_balances[bank] = balances


def read_folder(folder):
    """Read each source independently: a missing, duplicated or unreadable file blocks only what depends on it."""
    folder = Path(folder)
    bundle = ClosingInputs()
    for key, pattern in SOURCE_PATTERNS:
        matches = sorted(folder.glob(pattern))
        if len(matches) != 1:
            bundle.blocked[key] = {'arquivo':pattern, 'motivo':f'esperado 1 arquivo {pattern}; encontrados {len(matches)}'}
            bundle.block_causes[key] = 'AUSENTE_OU_DUPLICADO'
            continue
        path = matches[0]
        bundle.sources[path.name] = {'sha256':hashlib.sha256(path.read_bytes()).hexdigest(), 'path':str(path.resolve()), 'kind':key}
        try:
            load_source(bundle, key, path)
        except Exception as error:  # any reader failure becomes an explicit block, never a silent gap
            bundle.blocked[key] = {'arquivo':path.name, 'motivo':f'arquivo ilegível ou com layout diferente do esperado ({type(error).__name__}: {error})'}
            bundle.block_causes[key] = 'ILEGIVEL_OU_LAYOUT'
            bundle.sources[path.name]['status'] = 'BLOQUEADO'
    # Late Orbe adjustments (06_*.csv) are optional: same layout as 01, merged into the checkout rows.
    for path in sorted(folder.glob('06_*.csv')):
        bundle.sources[path.name] = {'sha256':hashlib.sha256(path.read_bytes()).hexdigest(), 'path':str(path.resolve()), 'kind':'checkout_adjustments'}
        try:
            bundle.checkout.extend(read_checkout(path))
        except Exception as error:
            previous = bundle.blocked.get('checkout_adjustments')
            reason = f'{path.name}: arquivo ilegível ou com layout diferente do esperado ({type(error).__name__}: {error})'
            bundle.block_causes['checkout_adjustments'] = 'ILEGIVEL_OU_LAYOUT'
            bundle.blocked['checkout_adjustments'] = {'arquivo':previous['arquivo']+', '+path.name if previous else path.name, 'motivo':previous['motivo']+'; '+reason if previous else reason}
            bundle.sources[path.name]['status'] = 'BLOQUEADO'
    if not bundle.sources:
        raise ValueError(f'No source files found in {folder}; check --input')
    return bundle


def read_inputs(path):
    path = Path(path)
    if path.is_dir():
        return read_folder(path)
    if path.suffix.lower() != '.zip' or not path.is_file():
        raise ValueError('Input must be a source folder or ZIP archive')
    try:
        with TemporaryDirectory(prefix='closing-input-') as directory, ZipFile(path) as archive:
            folder = Path(directory)
            names = set()
            if sum(item.file_size for item in archive.infolist()) > 100*1024*1024:
                raise ValueError('Archive exceeds the 100 MiB expanded-size limit')
            for item in archive.infolist():
                if item.is_dir():
                    continue
                if Path(item.filename).name != item.filename or '/' in item.filename or '\\' in item.filename or item.filename in names:
                    raise ValueError('Archive must use unique, flat filenames')
                names.add(item.filename)
                (folder/item.filename).write_bytes(archive.read(item))
            bundle = read_folder(folder)
            for name, metadata in bundle.sources.items():
                if 'path' in metadata:
                    bundle.attachments[name] = (folder/name).read_bytes()
                    del metadata['path']
                    metadata['parent'] = path.name
            bundle.sources[path.name] = {'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'path':str(path.resolve()),'kind':'archive'}
            return bundle
    except BadZipFile as error:
        raise ValueError(f'Invalid ZIP input: {path.name}') from error
