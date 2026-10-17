"""Deterministic validation of inference proposals (ADR 0002): schema, action/parameter pairing and verifiable evidence.

The model only proposes; nothing here trusts its text. Evidence is re-read from the source file.
"""
from io import BytesIO
from pathlib import Path
import json
import re

from openpyxl import load_workbook

from .privacy import mask_personal_data
from .readers import decode

ROOT = Path(__file__).resolve().parents[2]
_SCHEMA = {}


def load_schema(path=None):
    path = Path(path) if path else ROOT/'config'/'proposal.schema.json'
    key = str(path)
    if key not in _SCHEMA:
        _SCHEMA[key] = json.loads(path.read_text(encoding='utf-8'))
    return _SCHEMA[key]


def _derive_actions(schema):
    """{class: {action: [parameter keys]}}. ENCAMINHAR_HUMANO never takes parameters; the other actions take the
    non-empty parameter variants in order (the pairing itself is enforced by check_shape)."""
    actions = {}
    for variant in schema['anyOf']:
        props = variant['properties']
        params = props['parametros']
        options = params['anyOf'] if 'anyOf' in params else [params]
        keyed = [list(option['properties']) for option in options if option['properties']]
        mapping = {}
        for action in props['acao']['enum']:
            mapping[action] = [] if action == 'ENCAMINHAR_HUMANO' else (keyed.pop(0) if keyed else [])
        actions[props['classe']['const']] = mapping
    return actions


ACTIONS = _derive_actions(load_schema())

_TYPES = {'object':dict, 'array':list, 'string':str, 'boolean':bool}


def _check(value, schema, path, errors):
    if 'anyOf' in schema:
        attempts = []
        for option in schema['anyOf']:
            found = []
            _check(value, option, path, found)
            if not found:
                return
            attempts.append(found)
        best = min(attempts, key=len)
        errors.append(f'{path or "$"}: nenhuma variante do schema aceita o valor ({"; ".join(best)})')
        return
    kind = schema.get('type')
    if kind and not (isinstance(value, _TYPES[kind]) and not (kind != 'boolean' and isinstance(value, bool))):
        errors.append(f'{path or "$"}: esperado {kind}')
        return
    if 'const' in schema and value != schema['const']:
        errors.append(f'{path or "$"}: deve ser {schema["const"]!r}')
    if 'enum' in schema and value not in schema['enum']:
        errors.append(f'{path or "$"}: {value!r} fora de {schema["enum"]}')
    if isinstance(value, dict):
        for name in schema.get('required', []):
            if name not in value:
                errors.append(f'{path or "$"}: falta a propriedade {name}')
        properties = schema.get('properties', {})
        for name, item in value.items():
            if name in properties:
                _check(item, properties[name], f'{path}.{name}' if path else name, errors)
            elif schema.get('additionalProperties') is False:
                errors.append(f'{path or "$"}: propriedade não permitida {name}')
    if isinstance(value, list):
        if len(value) < schema.get('minItems', 0):
            errors.append(f'{path or "$"}: mínimo de {schema["minItems"]} item(ns)')
        if 'items' in schema:
            for index, item in enumerate(value):
                _check(item, schema['items'], f'{path}[{index}]', errors)


def check_schema(instance, schema):
    errors = []
    _check(instance, schema, '', errors)
    return errors


def _as_dict(excecao):
    return excecao.as_dict() if hasattr(excecao, 'as_dict') else excecao


def check_shape(proposal, excecao):
    excecao = _as_dict(excecao)
    errors = check_schema(proposal, load_schema())
    if errors:
        return errors
    if proposal['excecao_id'] != excecao['excecao_id']:
        errors.append(f'excecao_id {proposal["excecao_id"]!r} difere da exceção {excecao["excecao_id"]!r}')
    if proposal['classe'] != excecao['classe']:
        errors.append(f'classe {proposal["classe"]!r} difere da exceção {excecao["classe"]!r}')
    expected = ACTIONS.get(proposal['classe'], {}).get(proposal['acao'])
    if expected is None:
        errors.append(f'ação {proposal["acao"]!r} não existe para a classe {proposal["classe"]}')
    elif sorted(proposal['parametros']) != sorted(expected):
        errors.append(f'parametros {sorted(proposal["parametros"])} não correspondem à ação {proposal["acao"]} (esperado {sorted(expected)})')
    return errors


def _lines(content):
    # same line counting as csv.reader.line_num and the OFX reader: \r\n, \r and \n each end one line
    return re.split(r'\r\n|\r|\n', decode(content))


def _read_source(bundle, name):
    metadata = bundle.sources.get(name)
    if metadata and metadata.get('path'):
        path = Path(metadata['path'])
        if not path.is_file():
            raise ValueError(f'Source file not found: {name}')
        return path.read_bytes()
    if name in bundle.attachments:
        return bundle.attachments[name]
    raise ValueError(f'Source not available: {name}')


def source_line(bundle, ref):
    name, line = ref['arquivo'], ref['linha']
    content = _read_source(bundle, name)
    if name.lower().endswith('.xlsx'):
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
        try:
            if ref['aba'] not in workbook.sheetnames:
                raise ValueError(f'Sheet not found: {name}!{ref["aba"]}')
            rows = list(workbook[ref['aba']].iter_rows(min_row=line, max_row=line, values_only=True))
        finally:
            workbook.close()
        if not rows:
            raise ValueError(f'Row not found: {name}!{ref["aba"]} row {line}')
        return ';'.join('' if cell is None else str(cell) for cell in rows[0])
    lines = _lines(content)
    if not 1 <= line <= len(lines):
        raise ValueError(f'Line not found: {name} line {line}')
    if name.lower().endswith('.ofx'):
        block = []
        for text in lines[line-1:]:
            block.append(text)
            if '</STMTTRN>' in text:
                break
        return '\n'.join(block)
    return lines[line-1]


def check_evidence(proposal, excecao, bundle):
    excecao = _as_dict(excecao)
    refs = {}
    for r in excecao['refs']:
        refs.setdefault((r['arquivo'], r['ref']), []).append(r)
    errors = []
    for index, evidence in enumerate(proposal.get('evidencias') or []):
        label = f'evidencias[{index}]'
        if not isinstance(evidence, dict):
            errors.append(f'{label}: evidência malformada')
            continue
        candidates = refs.get((evidence.get('arquivo'), evidence.get('ref')))
        if not candidates:
            errors.append(f'{label}: {evidence.get("arquivo")!r} {evidence.get("ref")!r} não é uma referência da exceção')
            continue
        snippet = evidence.get('trecho')
        if not isinstance(snippet, str) or not snippet.strip():
            errors.append(f'{label}: trecho vazio')
            continue
        # the evidence has no sheet field: it holds if the snippet is in the line of any candidate ref (same file and ref)
        found, failures = False, []
        for ref in candidates:
            try:
                line = source_line(bundle, ref)
            except ValueError as error:
                failures.append(str(error))
                continue
            if snippet in line or snippet in mask_personal_data(line):
                found = True
                break
        if not found:
            errors.append(f'{label}: ' + ('; '.join(failures) if failures else 'trecho não existe literalmente na linha citada'))
    if not proposal.get('evidencias'):
        errors.append('evidencias: nenhuma evidência')
    return errors


def _result(check, errors, ok_reason):
    return {'check':check, 'status':'FAIL' if errors else 'PASS', 'reason':'; '.join(errors) if errors else ok_reason}


class Validator:
    classe = None

    def class_checks(self, proposal, excecao, bundle, policy):
        return []

    def validate(self, proposal, excecao, bundle, policy):
        checks = []
        for name, errors, ok in (('schema', lambda: check_shape(proposal, excecao), 'Forma conforme o schema e a exceção.'),
                                 ('evidencia', lambda: check_evidence(proposal, excecao, bundle), 'Trechos existem literalmente nas linhas citadas.')):
            checks.append(_result(name, errors(), ok))
            if checks[-1]['status'] == 'FAIL':
                return {'ok':False, 'checks':checks}
        for item in self.class_checks(proposal, excecao, bundle, policy):
            checks.append(item)
            if item['status'] == 'FAIL':
                break
        return {'ok':all(c['status'] == 'PASS' for c in checks), 'checks':checks}


VALIDATORS = {}


def register(cls):
    VALIDATORS[cls.classe] = cls
    return cls


def validator_for(classe):
    cls = VALIDATORS.get(classe)
    return cls() if cls else None


@register
class MovimentoSemRegraValidator(Validator):
    classe = 'MOVIMENTO_SEM_REGRA'

    def class_checks(self, proposal, excecao, bundle, policy):
        return [{'check':'validador_da_classe', 'status':'FAIL', 'reason':'Validador de MOVIMENTO_SEM_REGRA chega na Fase 3; até lá toda proposta vai para humano.'}]
